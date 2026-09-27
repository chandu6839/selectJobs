"""
Excel Agent
-----------
Goal: write each analysed job + ChatGPT's verdict to job_matches.xlsx.

Creates the workbook with headers if it doesn't exist yet, otherwise
appends a new row. No browser involved - pure file I/O via openpyxl.
"""

from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter

import config

HEADERS = [
    "Date",
    "Country",
    "Title",
    "Company",
    "URL",
    "Score",
    "Match",
    "Matching Keywords",
    "Reason",
    "Saved",
]


def _ensure_headers_up_to_date(ws):
    """
    If the workbook was created before a header was added (e.g. this file
    existed before the "Saved" column), append any missing headers as new
    trailing columns rather than silently misaligning old rows.
    """
    existing_headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []

    for header in HEADERS:
        if header not in existing_headers:
            new_col_idx = ws.max_column + 1
            ws.cell(row=1, column=new_col_idx, value=header)
            ws.column_dimensions[get_column_letter(new_col_idx)].width = max(15, len(header) + 4)


def _get_or_create_workbook():
    if config.EXCEL_OUTPUT_PATH.exists():
        wb = openpyxl.load_workbook(config.EXCEL_OUTPUT_PATH)
        ws = wb.active
        _ensure_headers_up_to_date(ws)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Job Matches"
        ws.append(HEADERS)
        for col_idx, header in enumerate(HEADERS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(15, len(header) + 4)
    return wb, ws


def run(job_data: dict, verdict: dict, saved_status: str = "", country: str = "") -> str:
    """
    job_data: dict with title/company/url/job_description (from LinkedIn Agent)
    verdict: dict with match/score/matching_keywords/reason (from ChatGPT Agent)
    saved_status: what happened to the job - "Queued" (matched, handed to
        CV generation; "+ saved on LinkedIn" if config.SAVE_ON_LINKEDIN),
        "No" (didn't qualify), "Skipped (company excluded)", or another
        "Skipped (...)" reason decided from the search-result card.
        Older rows may say "Yes"/"Failed" (from before the queue existed).
    country: which country this run searched (e.g. "Germany", "Ireland").

    Returns the path to the Excel file.
    """
    wb, ws = _get_or_create_workbook()

    # Write row values by header name (not fixed position), so it stays
    # correct even for files migrated from an older header layout.
    header_row = [cell.value for cell in ws[1]]
    row_values = {
        "Date": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Country": country,
        "Title": job_data.get("title", ""),
        "Company": job_data.get("company", ""),
        "URL": job_data.get("url", ""),
        "Score": verdict.get("score", ""),
        "Match": verdict.get("match", ""),
        "Matching Keywords": ", ".join(verdict.get("matching_keywords", [])),
        "Reason": verdict.get("reason", ""),
        "Saved": saved_status,
    }
    row = [row_values.get(header, "") for header in header_row]
    ws.append(row)

    wb.save(config.EXCEL_OUTPUT_PATH)
    print(f"[Excel Agent] Row written to {config.EXCEL_OUTPUT_PATH}")
    return str(config.EXCEL_OUTPUT_PATH)
